"""Tests for scalar batch evaluation and export."""

from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from forestagent.backends.mock_backend import MockBackend
from forestagent.evaluation import (
    evaluate_scalar_tasks,
    export_evaluation_csv,
    export_evaluation_xlsx,
)


class EvaluationTests(unittest.TestCase):
    def test_evaluate_scalar_tasks_computes_metrics(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir)
            self._create_sample_dataset(
                data_dir,
                rows=[(150, "银杏", "150_64.las", 1, 1.0, 2.0, 3.0, 10.0, 8.0, 2.0, 4.0)],
                ground_files=["150_64.las"],
                air_files=["150_64.las"],
            )

            result = evaluate_scalar_tasks(
                data_dir=data_dir,
                task_ids=["q1_dbh"],
                sample_ids=["150_64"],
                modalities=["ground", "air"],
                backend=MockBackend(),
            )

            self.assertEqual(len(result.records), 2)
            self.assertTrue(all(record.status == "success" for record in result.records))

            ground_metrics = next(
                metric
                for metric in result.metrics
                if metric.task_id == "q1_dbh" and metric.modality == "ground"
            )
            self.assertEqual(ground_metrics.total_records, 1)
            self.assertEqual(ground_metrics.success_count, 1)
            self.assertEqual(ground_metrics.success_rate, 1.0)
            self.assertAlmostEqual(ground_metrics.mae, 18.4, places=6)
            self.assertAlmostEqual(ground_metrics.rmse, 18.4, places=6)
            self.assertAlmostEqual(ground_metrics.mean_bias, 18.4, places=6)

            comparison = next(
                item for item in result.modality_comparison if item.task_id == "q1_dbh"
            )
            self.assertEqual(comparison.ground_mae, 18.4)
            self.assertEqual(comparison.air_mae, 18.4)

    def test_evaluate_scalar_tasks_handles_missing_modality_missing_measured_and_failure(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir)
            self._create_sample_dataset(
                data_dir,
                rows=[(151, "银杏", "151_01.las", 2, 4.0, 5.0, 6.0, 20.0, None, None, None)],
                ground_files=["151_01.las"],
                air_files=["999_01.las"],
            )

            result = evaluate_scalar_tasks(
                data_dir=data_dir,
                task_ids=["q1_dbh", "q2_height", "q3_crown_width"],
                sample_ids=["151_01"],
                modalities=["ground", "air"],
                backend=MockBackend(failures={"estimate_height": "Height failure."}),
            )

            by_key = {(record.task_id, record.modality): record for record in result.records}

            self.assertEqual(by_key[("q1_dbh", "ground")].status, "success")
            self.assertEqual(by_key[("q2_height", "ground")].status, "task_failed")
            self.assertEqual(by_key[("q3_crown_width", "ground")].status, "missing_measured")
            self.assertEqual(by_key[("q1_dbh", "air")].status, "missing_modality")

            q2_ground_metrics = next(
                metric
                for metric in result.metrics
                if metric.task_id == "q2_height" and metric.modality == "ground"
            )
            self.assertEqual(q2_ground_metrics.success_count, 0)
            self.assertEqual(q2_ground_metrics.success_rate, 0.0)
            self.assertIsNone(q2_ground_metrics.mae)

    def test_export_evaluation_csv_and_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir) / "data"
            output_dir = Path(temp_dir) / "exports"
            self._create_sample_dataset(
                data_dir,
                rows=[(150, "银杏", "150_64.las", 1, 1.0, 2.0, 3.0, 10.0, 8.0, 2.0, 4.0)],
                ground_files=["150_64.las"],
                air_files=["150_64.las"],
            )

            result = evaluate_scalar_tasks(
                data_dir=data_dir,
                task_ids=["q1_dbh"],
                sample_ids=["150_64"],
                modalities=["ground"],
                backend=MockBackend(),
            )

            csv_path = export_evaluation_csv(result, output_dir / "records.csv")
            xlsx_path = export_evaluation_xlsx(result, output_dir / "records.xlsx")

            self.assertTrue(csv_path.exists())
            self.assertTrue(xlsx_path.exists())

            with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                rows = list(reader)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["sample_id"], "150_64")
            self.assertEqual(rows[0]["task_id"], "q1_dbh")

            workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    ["records", "metrics", "modality_comparison"],
                )
                records_sheet = workbook["records"]
                self.assertEqual(records_sheet.max_row, 2)
            finally:
                workbook.close()

    @staticmethod
    def _create_sample_dataset(
        data_dir: Path,
        rows: list[tuple[object, ...]],
        ground_files: list[str],
        air_files: list[str],
    ) -> None:
        ground_dir = data_dir / "ground_points"
        air_dir = data_dir / "air_points"
        for row in rows:
            plot_id = str(row[0])
            (ground_dir / plot_id).mkdir(parents=True, exist_ok=True)
            (air_dir / plot_id).mkdir(parents=True, exist_ok=True)

        for file_name in ground_files:
            plot_id = file_name.split("_", 1)[0]
            (ground_dir / plot_id).mkdir(parents=True, exist_ok=True)
            (ground_dir / plot_id / file_name).write_text("", encoding="utf-8")
        for file_name in air_files:
            plot_id = file_name.split("_", 1)[0]
            (air_dir / plot_id).mkdir(parents=True, exist_ok=True)
            (air_dir / plot_id / file_name).write_text("", encoding="utf-8")

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet2"
        sheet.append(
            [
                "样地号",
                "树种",
                "对应的文件名",
                "编号",
                "X",
                "Y",
                "Z",
                "胸径cm",
                "树高m",
                "东西冠幅m",
                "南北冠幅m",
            ]
        )
        for row in rows:
            sheet.append(list(row))
        workbook.save(data_dir / "reference.xlsx")
        workbook.close()


if __name__ == "__main__":
    unittest.main()
