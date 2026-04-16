"""Dataset catalog loading for real point cloud samples and reference tables."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field


class MeasuredTreeAttributes(BaseModel):
    """Measured reference attributes loaded from the Excel table."""

    model_config = ConfigDict(extra="forbid")

    species: str
    x: float
    y: float
    z: float
    dbh_cm: float | None = None
    height_m: float | None = None
    crown_width_east_west_m: float | None = None
    crown_width_north_south_m: float | None = None

    @property
    def crown_width_mean_m(self) -> float | None:
        """Arithmetic mean of east-west and north-south measured crown widths."""

        if (
            self.crown_width_east_west_m is None
            or self.crown_width_north_south_m is None
        ):
            return None
        return (self.crown_width_east_west_m + self.crown_width_north_south_m) / 2.0


class TreeSampleRecord(BaseModel):
    """One tree-level sample with optional ground/air point cloud paths."""

    model_config = ConfigDict(extra="forbid")

    sample_id: str = Field(min_length=1)
    plot_id: str = Field(min_length=1)
    file_name: str = Field(min_length=1)
    tree_number: int
    measured: MeasuredTreeAttributes
    ground_point_cloud_path: str | None = None
    air_point_cloud_path: str | None = None

    def available_modalities(self) -> list[str]:
        """Return modalities that have a point cloud file on disk."""

        modalities: list[str] = []
        if self.ground_point_cloud_path:
            modalities.append("ground")
        if self.air_point_cloud_path:
            modalities.append("air")
        return modalities

    def point_cloud_path_for(self, modality: str) -> str:
        """Return the file path for the requested modality."""

        if modality == "ground" and self.ground_point_cloud_path:
            return self.ground_point_cloud_path
        if modality == "air" and self.air_point_cloud_path:
            return self.air_point_cloud_path
        raise FileNotFoundError(
            f"Sample {self.sample_id} does not have a {modality!r} point cloud."
        )

    def measured_summary(self) -> dict[str, float | str | None]:
        """Return a compact measured reference summary for downstream reporting."""

        return {
            "species": self.measured.species,
            "dbh_cm": self.measured.dbh_cm,
            "height_m": self.measured.height_m,
            "crown_width_east_west_m": self.measured.crown_width_east_west_m,
            "crown_width_north_south_m": self.measured.crown_width_north_south_m,
            "crown_width_mean_m": self.measured.crown_width_mean_m,
            "x": self.measured.x,
            "y": self.measured.y,
            "z": self.measured.z,
        }


class DataCatalog:
    """Load and serve tree-level records from the data directory."""

    _EXPECTED_HEADERS = [
        "样地号",
        "树种",
        "对应的文件名",
        "编号",
        "X",
        "Y",
        "Z",
        "胸径cm",
        "树高m",
        "东西冠幅m",
        "南北冠幅m",
    ]

    def __init__(self, records: dict[str, TreeSampleRecord], data_dir: str | Path) -> None:
        self._records = dict(sorted(records.items()))
        self.data_dir = str(Path(data_dir).resolve())

    @classmethod
    def from_data_dir(cls, data_dir: str | Path) -> "DataCatalog":
        """Load the catalog from a data directory containing LAS files and one workbook."""

        data_path = Path(data_dir).resolve()
        workbook_path = cls._find_reference_workbook(data_path)
        modality_roots = cls._discover_modality_roots(data_path)
        records = cls._load_records(workbook_path, modality_roots)
        return cls(records=records, data_dir=data_path)

    def get_record(self, sample_id: str) -> TreeSampleRecord:
        """Fetch one sample record by sample id."""

        if sample_id not in self._records:
            raise KeyError(f"Unknown sample_id: {sample_id}")
        return self._records[sample_id]

    def list_records(self, limit: int | None = None) -> list[TreeSampleRecord]:
        """Return records in stable sample-id order."""

        records = list(self._records.values())
        if limit is None:
            return records
        return records[:limit]

    def summary(self) -> dict[str, int]:
        """Return a compact summary of the loaded catalog."""

        ground_count = 0
        air_count = 0
        dual_modality_count = 0
        for record in self._records.values():
            modalities = record.available_modalities()
            if "ground" in modalities:
                ground_count += 1
            if "air" in modalities:
                air_count += 1
            if len(modalities) == 2:
                dual_modality_count += 1

        return {
            "record_count": len(self._records),
            "ground_count": ground_count,
            "air_count": air_count,
            "dual_modality_count": dual_modality_count,
        }

    @classmethod
    def _find_reference_workbook(cls, data_path: Path) -> Path:
        workbooks = sorted(data_path.glob("*.xlsx"))
        if not workbooks:
            raise FileNotFoundError(f"No workbook found under {data_path}.")
        if len(workbooks) > 1:
            raise ValueError(f"Expected one workbook under {data_path}, found {len(workbooks)}.")
        return workbooks[0]

    @classmethod
    def _discover_modality_roots(cls, data_path: Path) -> dict[str, Path]:
        modality_roots: dict[str, Path] = {}
        for child in data_path.iterdir():
            if not child.is_dir():
                continue
            if not any(child.rglob("*.las")):
                continue
            try:
                modality = cls._infer_modality(child.name)
            except ValueError:
                continue
            modality_roots[modality] = child

        if "ground" not in modality_roots:
            raise FileNotFoundError("Ground point cloud directory was not found.")
        if "air" not in modality_roots:
            raise FileNotFoundError("Air point cloud directory was not found.")
        return modality_roots

    @classmethod
    def _infer_modality(cls, directory_name: str) -> str:
        normalized = directory_name.lower()
        if "ground" in normalized or "tls" in normalized or "地基" in directory_name:
            return "ground"
        if "air" in normalized or "uav" in normalized or "空基" in directory_name:
            return "air"
        if "uls" in normalized:
            return "air"
        raise ValueError(f"Could not infer modality from directory name: {directory_name}")

    @classmethod
    def _load_records(
        cls, workbook_path: Path, modality_roots: dict[str, Path]
    ) -> dict[str, TreeSampleRecord]:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            worksheet = workbook[workbook.sheetnames[0]]

            header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
            headers = [str(value) if value is not None else "" for value in header_row]
            if headers != cls._EXPECTED_HEADERS:
                raise ValueError(f"Unexpected workbook headers: {headers!r}")

            records: dict[str, TreeSampleRecord] = {}
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if row[0] is None or row[2] is None:
                    continue

                plot_id = str(row[0])
                file_name = str(row[2])
                sample_id = Path(file_name).stem
                if sample_id in records:
                    raise ValueError(f"Duplicate sample_id detected: {sample_id}")

                ground_path = modality_roots["ground"] / plot_id / file_name
                air_path = modality_roots["air"] / plot_id / file_name

                records[sample_id] = TreeSampleRecord(
                    sample_id=sample_id,
                    plot_id=plot_id,
                    file_name=file_name,
                    tree_number=int(row[3]),
                    measured=MeasuredTreeAttributes(
                        species=str(row[1]),
                        x=float(row[4]),
                        y=float(row[5]),
                        z=float(row[6]),
                        dbh_cm=cls._to_optional_float(row[7]),
                        height_m=cls._to_optional_float(row[8]),
                        crown_width_east_west_m=cls._to_optional_float(row[9]),
                        crown_width_north_south_m=cls._to_optional_float(row[10]),
                    ),
                    ground_point_cloud_path=str(ground_path.resolve())
                    if ground_path.exists()
                    else None,
                    air_point_cloud_path=str(air_path.resolve()) if air_path.exists() else None,
                )

            return records
        finally:
            workbook.close()

    @staticmethod
    def _to_optional_float(value: object) -> float | None:
        if value is None:
            return None
        if isinstance(value, str) and not value.strip():
            return None
        return float(value)
