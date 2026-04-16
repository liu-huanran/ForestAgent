"""Export helpers for batch evaluation outputs."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook

from .batch_evaluator import BatchEvaluationResult


def export_evaluation_csv(result: BatchEvaluationResult, output_path: str | Path) -> Path:
    """Export per-sample evaluation records to CSV."""

    path = Path(output_path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    records = [record.model_dump() for record in result.records]
    fieldnames = list(records[0].keys()) if records else []

    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in records:
            writer.writerow(row)

    return path


def export_evaluation_xlsx(result: BatchEvaluationResult, output_path: str | Path) -> Path:
    """Export records, metrics, and modality comparison to XLSX."""

    path = Path(output_path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    try:
        _write_sheet(
            workbook.active,
            "records",
            [record.model_dump() for record in result.records],
        )
        _write_sheet(
            workbook.create_sheet("metrics"),
            "metrics",
            [metric.model_dump() for metric in result.metrics],
        )
        _write_sheet(
            workbook.create_sheet("modality_comparison"),
            "modality_comparison",
            [item.model_dump() for item in result.modality_comparison],
        )
        workbook.save(path)
    finally:
        workbook.close()

    return path


def _write_sheet(worksheet, title: str, rows: list[dict[str, object]]) -> None:
    worksheet.title = title
    if not rows:
        return

    headers = list(rows[0].keys())
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header) for header in headers])

